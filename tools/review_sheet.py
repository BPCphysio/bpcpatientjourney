"""Write the reference answers out as a workbook the senior physiotherapist can
review, and mark up, without touching JSON.

    python tools/review_sheet.py [output.xlsx]

Three sheets: how to review, the key points (what earns marks), and the
reference answers (ten full-marks answers per question). Yellow cells are
for the reviewer. Send the marked-up file back and the corrections go into
grading/answers.json.
"""
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
OUT = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / 'grading' / 'answers-review.xlsx'

ref = json.loads((ROOT / 'grading' / 'answers.json').read_text(encoding='utf-8'))
titles = json.loads((ROOT / 'grading' / 'titles.json').read_text(encoding='utf-8'))

FONT = 'Arial'
HEAD = Font(name=FONT, bold=True, color='FFFFFF')
HEAD_FILL = PatternFill('solid', fgColor='2F4F4F')
FILL_IN = PatternFill('solid', fgColor='FFFF00')
BODY = Font(name=FONT, size=10)
BOLD = Font(name=FONT, size=10, bold=True)
WRAP = Alignment(wrap_text=True, vertical='top')
THIN = Side(style='thin', color='BFBFBF')
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
QNAME = {'q1': 'The three questions', 'q2': 'Why those three'}
STYLES = ['Phone questions', 'Terse notes', 'Clinical', 'Plain', 'Mixed']


def table(ws, headers, widths, rows, fill_cols):
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font, cell.fill, cell.alignment, cell.border = HEAD, HEAD_FILL, WRAP, BOX
    for r in rows:
        ws.append(r)
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            cell.font, cell.alignment, cell.border = BODY, WRAP, BOX
            if cell.column in fill_cols:
                cell.fill = FILL_IN
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions


wb = Workbook()

# ---------------------------------------------------------------- how to
how = wb.active
how.title = 'How to review'
lines = [
    ('Reference answers for the auto-grader', BOLD),
    ('', BODY),
    ('The auto-grader marks the two written answers in every case: "the three most important questions" '
     'and "why those three". Each has three key points. An answer earns a point if it uses one of the '
     'point\'s cue words, or if it reads like one of the ten reference answers for that point. '
     'The percentage is the share of points earned; the suggested mark out of 5 is the percentage divided by 20.', BODY),
    ('', BODY),
    ('What to check', BOLD),
    ('Key points sheet: is each point right for the case, and are the cue words ones that really mean '
     'this point? A cue word that fits every case (for example "pain") would hand out marks too easily.', BODY),
    ('Reference answers sheet: would you give each of these answers full marks? All ten per question are '
     'meant to be full-marks answers, written the ways different physiotherapists write.', BODY),
    ('', BODY),
    ('Which cells to fill in', BOLD),
    ('Only the yellow columns. Leave everything else as it is.', BODY),
    ('  OK? — write "yes" or "no".', BODY),
    ('  Change to / Comment — when you write "no", say what it should be instead.', BODY),
    ('', BODY),
    ('Example of a filled-in row (Key points sheet)', BOLD),
    ('  c01 · The three questions · point 3 · OK? no · Change to: also count "breastfeeding" and "ให้นมบุตร" as cue words', BODY),
    ('', BODY),
    ('Status: ' + ref.get('_status', ''), BODY),
    ('Send the file back when done; the corrections go straight into the grader.', BODY),
]
for text, font in lines:
    how.append([text])
    how.cell(row=how.max_row, column=1).font = font
    how.cell(row=how.max_row, column=1).alignment = WRAP
how.column_dimensions['A'].width = 120

# ---------------------------------------------------------------- points
pts = wb.create_sheet('Key points')
rows = []
for cid, case in ref['cases'].items():
    t = titles.get(cid, {})
    for q in ('q1', 'q2'):
        for i, p in enumerate(case[q]['points'], 1):
            rows.append([cid, t.get('en', ''), t.get('th', ''), QNAME[q], i, p['en'], p['th'],
                         ', '.join(p.get('core', [])), ', '.join(p.get('near', [])), '', ''])
table(pts, ['Case', 'Title', 'ชื่อเคส', 'Question', 'Point', 'Key point (English)', 'ประเด็นหลัก (ไทย)',
            'Cue words — full credit', 'Cue words — half credit', 'OK?', 'Change to'],
      [7, 22, 20, 18, 7, 34, 40, 44, 30, 8, 36], rows, fill_cols={10, 11})

# ---------------------------------------------------------------- answers
ans = wb.create_sheet('Reference answers')
rows = []
for cid, case in ref['cases'].items():
    t = titles.get(cid, {})
    for q in ('q1', 'q2'):
        for lang in ('th', 'en'):
            for n, a in enumerate(case[q]['answers'][lang]):
                rows.append([cid, t.get('th' if lang == 'th' else 'en', ''), QNAME[q],
                             'ไทย' if lang == 'th' else 'English', n + 1, STYLES[n] if n < 5 else '',
                             a[0], a[1], a[2], '', ''])
table(ans, ['Case', 'Title', 'Question', 'Language', '#', 'Style', 'Point 1', 'Point 2', 'Point 3', 'OK?', 'Comment'],
      [7, 22, 18, 9, 4, 14, 44, 44, 44, 8, 34], rows, fill_cols={10, 11})

OUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUT)
print(f'wrote {OUT} — {pts.max_row - 1} key points, {ans.max_row - 1} reference answers')
